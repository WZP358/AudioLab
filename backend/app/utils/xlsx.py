"""轻量 .xlsx 导出工具：把表头 + 行数据写成 Excel 字节流（UTF-8 中文友好）。"""
from __future__ import annotations

import io
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def build_xlsx(
    headers: Sequence[str],
    rows: Iterable[Sequence],
    sheet_title: str = "Sheet1",
    freeze_header: bool = True,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel 工作表名上限 31 字符
    ws.append(list(headers))
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0891B2")
    for idx, _cell in enumerate(ws[1], start=1):
        _cell.font = header_font
        _cell.fill = header_fill
    for row in rows:
        ws.append([_xlsx_value(v) for v in row])
    if freeze_header:
        ws.freeze_panes = "A2"
    # 自动列宽（按表头 + 抽样前 100 行估算，避免过宽）
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1]))
        for row in ws.iter_rows(min_row=2, max_row=min(101, ws.max_row), min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, min(60, len(str(cell.value))))
        ws.column_dimensions[letter].width = max(8, max_len + 2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_value(v):
    # openpyxl 不能写 numpy 标量/复杂对象，统一转基本类型
    if v is None:
        return None
    if isinstance(v, (str, int, float, bool)):
        return v
    try:
        if hasattr(v, "item"):
            v = v.item()
        return float(v) if isinstance(v, (int, float)) else str(v)
    except Exception:
        return str(v)
